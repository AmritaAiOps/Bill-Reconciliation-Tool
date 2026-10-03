"""Tests for the live estimate parser (app.pdf_estimate), the Excel column
layout / migration, and which PDFs get removed after a run.

Run from reconciliation_tool/:  python -m pytest tests -q
"""
import sys
import types

from openpyxl import Workbook, load_workbook
import pytest

try:
    import webview  # noqa: F401
except ImportError:  # the Api tests never open a window
    sys.modules["webview"] = types.ModuleType("webview")

from app import processing, web_ui
from app.excel_store import COLUMNS, ExcelStore
from app.pdf_bill import BillExtractionError
from app.pdf_estimate import parse_estimate_text

ESTIMATE_TEXT = """\
                            ESTIMATE     PERFORMA
     M R D Number : 317251                   Patient Name : ALAALDEN KHALIFA
                                                          OSMAN ALTAHER
     Admitting Doctor : Anil Kumar Murarka   Duration of Stay : 2
     Speciality   : Plastic-Reconstructive Surgery Expected D.O.A : 15/09/2026
     Bed Type     : IPS CatB_ V1_Semi Private Estimation : Estimation 2

    Estimated Expenditure:

     Particulars            Type               Quantity        Amount (Rs)

     Burn contracture release- Single Service  1.0             74360.00
     region
     Suturing- Moderate     Service            0.75            44550.00
     Medicines and Consumables                                 80000.00
     (APPROX)
     Investigations (APPROX)                                   20000.00
     Doctors Visits Charges                                    6400.00
     Medical Supervision Charges                               20000.00
     Bed                    Bed                2.0             9000.00
    Total Estimate                                             254310.00
"""

OLD_COLUMNS = [
    "MRD Number", "Patient Name", "Admitting Doctor", "Speciality", "Bed Type",
    "Duration of Stay", "Estimation", "Expected D.O.A",
    "Estimate Amount", "Billed Amount", "Difference", "Status",
]


# ---- Estimate parsing ------------------------------------------------------

def test_estimate_new_fields():
    r = parse_estimate_text(ESTIMATE_TEXT)
    assert r["admitting_doctor"] == "Anil Kumar Murarka"
    assert r["speciality"] == "Plastic-Reconstructive Surgery"
    assert r["surgery"] == "Burn contracture release- Single region; Suturing- Moderate"
    assert r["duration_days"] == 2
    assert r["estimate_amount"] == 254310.0


def test_duration_falls_back_to_header_without_bed_row():
    text = ESTIMATE_TEXT.replace(
        "     Bed                    Bed                2.0             9000.00\n", "")
    text = text.replace("Duration of Stay : 2", "Duration of Stay : 3")
    assert parse_estimate_text(text)["duration_days"] == 3


def test_bed_quantity_wins_over_header():
    text = ESTIMATE_TEXT.replace("Duration of Stay : 2", "Duration of Stay : 5")
    assert parse_estimate_text(text)["duration_days"] == 2


def test_no_service_rows_gives_empty_surgery():
    text = ESTIMATE_TEXT.replace("Single Service", "Single Package ").replace(
        "Moderate     Service", "Moderate     Package")
    assert parse_estimate_text(text)["surgery"] == ""


# ---- Excel store -----------------------------------------------------------

def test_add_estimate_writes_new_columns(tmp_path):
    store = ExcelStore(tmp_path / "out.xlsx")
    store.add_estimate(parse_estimate_text(ESTIMATE_TEXT))
    row = store.all_rows()[0]
    assert row["Doctor Name"] == "Anil Kumar Murarka"
    assert row["Speciality"] == "Plastic-Reconstructive Surgery"
    assert row["Surgery"] == "Burn contracture release- Single region; Suturing- Moderate"
    assert row["Duration of Stay (Days)"] == 2


def test_old_workbook_is_migrated(tmp_path):
    path = tmp_path / "old.xlsx"
    wb = Workbook()
    wb.active.append(OLD_COLUMNS)
    wb.active.append(["111", "JANE DOE", "Dr X", "Ortho", "Ward", "4", "Estimation 1",
                      "01/01/2026", 1000.0, 900.0, 100.0, "Less Amount"])
    wb.save(path)

    store = ExcelStore(path)
    store.save()
    assert [c.value for c in load_workbook(path).active[1]] == COLUMNS
    row = store.all_rows()[0]
    assert row["Doctor Name"] == "Dr X"
    assert row["Duration of Stay (Days)"] == "4"
    assert row["Surgery"] is None
    assert row["Billed Amount"] == 900.0 and row["Status"] == "Less Amount"
    assert store.has_mrd("111")


# ---- Which PDFs are marked done -------------------------------------------

def _pdfs(folder, *names):
    folder.mkdir()
    for n in names:
        (folder / n).write_bytes(b"%PDF")
    return folder


def test_estimate_done_files(tmp_path, monkeypatch):
    store = ExcelStore(tmp_path / "out.xlsx")
    store.add_estimate({"mrd_number": "2", "patient_name": "SAME PERSON"})
    store.add_estimate({"mrd_number": "3", "patient_name": "SOMEONE ELSE"})
    by_file = {
        "a_new.pdf": {"mrd_number": "1", "patient_name": "NEW"},
        "b_dup.pdf": {"mrd_number": "2", "patient_name": "same  person"},
        "c_conflict.pdf": {"mrd_number": "3", "patient_name": "DIFFERENT"},
    }
    monkeypatch.setattr(processing, "extract_estimate_fields", lambda f: dict(by_file[f.name]))
    folder = _pdfs(tmp_path / "est", *by_file)

    r = processing.process_estimate_folder(store, folder)
    assert sorted(f.name for f in r["done_files"]) == ["a_new.pdf", "b_dup.pdf"]


def test_bill_done_files(tmp_path, monkeypatch):
    store = ExcelStore(tmp_path / "out.xlsx")
    for mrd in ("1", "2", "3"):
        store.add_estimate({"mrd_number": mrd, "patient_name": "P", "estimate_amount": 100.0})
    ok = {"billed_amount": 100.0, "annexure_sanity_ok": True,
          "annexure_charge_sum": 100.0, "patient_name": "P"}
    by_file = {
        "a_clean.pdf": {**ok, "mrd_number": "1"},
        "b_name_mismatch.pdf": {**ok, "mrd_number": "2", "patient_name": "Q"},
        "c_bad_sum.pdf": {**ok, "mrd_number": "3", "annexure_sanity_ok": False,
                          "annexure_charge_sum": 50.0},
        "d_unmatched.pdf": {**ok, "mrd_number": "9"},
    }

    def fake(f):
        if f.name == "e_unreadable.pdf":
            raise BillExtractionError("nope")
        return dict(by_file[f.name])

    monkeypatch.setattr(processing, "extract_bill_fields", fake)
    folder = _pdfs(tmp_path / "bill", *by_file, "e_unreadable.pdf")

    r = processing.process_billing_folder(store, folder)
    assert [f.name for f in r["done_files"]] == ["a_clean.pdf"]


# ---- Removal only after a successful save ---------------------------------

@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(processing, "extract_estimate_fields",
                        lambda f: {"mrd_number": "1", "patient_name": "P"})
    recycled = []
    monkeypatch.setattr(web_ui, "recycle_files", lambda paths: recycled.extend(paths) or [])
    a = web_ui.Api({"output_dir": tmp_path / "out", "log_dir": tmp_path / "logs"})
    a.recycled = recycled
    return a


def test_files_recycled_after_save(api, tmp_path):
    folder = _pdfs(tmp_path / "est", "x.pdf")
    result = api._run(str(folder), None)
    assert result["ok"] and result["estimate"]["removed"] == 1
    assert [f.name for f in api.recycled] == ["x.pdf"]
    assert "done_files" not in result["estimate"]


def test_nothing_recycled_when_save_fails(api, tmp_path, monkeypatch):
    def locked():
        raise PermissionError
    monkeypatch.setattr(api._store, "save", locked)
    folder = _pdfs(tmp_path / "est", "x.pdf")
    assert api._run(str(folder), None)["ok"] is False
    assert api.recycled == []
