"""Excel-backed patient store: one row per patient, keyed on MRD Number."""
import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)

COLUMNS = [
    "MRD Number", "Patient Name", "Doctor Name", "Speciality", "Surgery", "Bed Type",
    "Duration of Stay (Days)", "Estimation", "Expected D.O.A",
    "Estimate Amount", "Billed Amount", "Difference", "Status",
]
COL_INDEX = {name: i + 1 for i, name in enumerate(COLUMNS)}

# Old header names -> current ones, for migrating workbooks made by earlier versions.
_HEADER_ALIASES = {
    "Admitting Doctor": "Doctor Name",
    "Duration of Stay": "Duration of Stay (Days)",
}

STATUS_CLOSED = "Closed"
STATUS_LESS = "Less Amount"
STATUS_EXCESS = "Excess Amount"
STATUS_PENDING = "Pending"

_ALL_STATUSES = (STATUS_CLOSED, STATUS_LESS, STATUS_EXCESS, STATUS_PENDING)


def normalize_name(name) -> str:
    """Case- and whitespace-insensitive form of a patient name, for comparison."""
    return " ".join(str(name or "").split()).lower()


def compute_status(estimate_amount, billed_amount):
    """Return (status, difference). difference is Estimate - Billed."""
    if billed_amount is None:
        return STATUS_PENDING, None
    diff = round((estimate_amount or 0) - billed_amount, 2)
    if abs(diff) < 0.005:
        return STATUS_CLOSED, diff
    if diff > 0:
        return STATUS_LESS, diff
    return STATUS_EXCESS, diff


class ExcelStore:
    """Wraps a single-sheet workbook with one row per patient (MRD Number key)."""

    def __init__(self, path):
        self.path = Path(path)
        if self.path.exists():
            self.wb = load_workbook(self.path)
            self.ws = self.wb.active
            self._migrate_columns()
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Reconciliation"
            self.ws.append(COLUMNS)
        self._index_rows()

    def _migrate_columns(self):
        """Rewrite the sheet under the current COLUMNS if its header row differs
        (cells are addressed by position, so an old layout would be misread)."""
        header = [c.value for c in self.ws[1]] if self.ws.max_row >= 1 else []
        while header and header[-1] is None:
            header.pop()
        if header == COLUMNS:
            return
        names = [_HEADER_ALIASES.get(h, h) for h in header]
        rows = []
        for values in self.ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in values):
                rows.append(dict(zip(names, values)))
        self.ws.delete_rows(1, self.ws.max_row)
        self.ws.append(COLUMNS)
        for r in rows:
            self.ws.append([r.get(name) for name in COLUMNS])
        dropped = [n for n in names if n and n not in COLUMNS]
        logger.info("Migrated %s to the current column layout (%d rows)%s", self.path.name,
                    len(rows), f"; dropped columns {dropped}" if dropped else "")

    def _index_rows(self):
        self.mrd_row = {}
        for r in range(2, self.ws.max_row + 1):
            mrd = self.ws.cell(row=r, column=COL_INDEX["MRD Number"]).value
            if mrd:
                self.mrd_row[str(mrd)] = r

    def has_mrd(self, mrd_number) -> bool:
        return str(mrd_number) in self.mrd_row

    def patient_name(self, mrd_number):
        r = self.mrd_row.get(str(mrd_number))
        return self.ws.cell(row=r, column=COL_INDEX["Patient Name"]).value if r else None

    def find_mrds_by_name(self, name) -> list:
        """MRD Numbers whose Patient Name matches `name` (normalized)."""
        target = normalize_name(name)
        if not target:
            return []
        return [mrd for mrd, r in self.mrd_row.items()
                if normalize_name(self.ws.cell(row=r, column=COL_INDEX["Patient Name"]).value) == target]

    def add_estimate(self, fields: dict) -> bool:
        """Add a new patient row from extracted estimate fields.
        Returns False (no-op) if the MRD Number is already present."""
        mrd = str(fields["mrd_number"])
        if mrd in self.mrd_row:
            return False
        row = [None] * len(COLUMNS)
        row[COL_INDEX["MRD Number"] - 1] = mrd
        row[COL_INDEX["Patient Name"] - 1] = fields.get("patient_name", "")
        row[COL_INDEX["Doctor Name"] - 1] = fields.get("admitting_doctor", "")
        row[COL_INDEX["Speciality"] - 1] = fields.get("speciality", "")
        row[COL_INDEX["Surgery"] - 1] = fields.get("surgery", "")
        row[COL_INDEX["Bed Type"] - 1] = fields.get("bed_type", "")
        row[COL_INDEX["Duration of Stay (Days)"] - 1] = fields.get(
            "duration_days", fields.get("duration_of_stay", ""))
        row[COL_INDEX["Estimation"] - 1] = fields.get("estimation_label", "")
        row[COL_INDEX["Expected D.O.A"] - 1] = fields.get("expected_doa", "")
        row[COL_INDEX["Estimate Amount"] - 1] = fields.get("estimate_amount")
        row[COL_INDEX["Status"] - 1] = STATUS_PENDING
        self.ws.append(row)
        self.mrd_row[mrd] = self.ws.max_row
        return True

    def apply_bill(self, fields: dict):
        """Match a bill's fields to an existing patient row by MRD Number and
        set Billed Amount / Difference / Status.

        Returns a result dict on success, or None if the MRD Number has no
        matching estimate row yet. The result's `name_mismatch` is True when
        the bill's patient name differs from the estimate row's; the bill is
        still applied (MRD is the match key) but the caller should flag it."""
        mrd = str(fields["mrd_number"])
        if mrd not in self.mrd_row:
            return None
        r = self.mrd_row[mrd]

        excel_name = self.ws.cell(row=r, column=COL_INDEX["Patient Name"]).value or ""
        bill_name = fields.get("patient_name", "")
        name_mismatch = bool(excel_name and bill_name
                             and normalize_name(excel_name) != normalize_name(bill_name))
        if name_mismatch:
            logger.warning(
                "Patient name mismatch for MRD %s: Excel has '%s', bill has '%s'",
                mrd, excel_name, bill_name,
            )

        estimate_amount = self.ws.cell(row=r, column=COL_INDEX["Estimate Amount"]).value
        billed_amount = fields["billed_amount"]
        status, diff = compute_status(estimate_amount, billed_amount)

        self.ws.cell(row=r, column=COL_INDEX["Billed Amount"], value=billed_amount)
        self.ws.cell(row=r, column=COL_INDEX["Difference"], value=diff)
        self.ws.cell(row=r, column=COL_INDEX["Status"], value=status)

        return {
            "mrd_number": mrd,
            "patient_name": excel_name or bill_name,
            "estimate_amount": estimate_amount,
            "billed_amount": billed_amount,
            "difference": diff,
            "status": status,
            "bill_patient_name": bill_name,
            "name_mismatch": name_mismatch,
        }

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(self.path)

    def overall_counts(self) -> dict:
        counts = {"Total Patients": 0, STATUS_CLOSED: 0, STATUS_LESS: 0, STATUS_EXCESS: 0, STATUS_PENDING: 0}
        for r in range(2, self.ws.max_row + 1):
            mrd = self.ws.cell(row=r, column=COL_INDEX["MRD Number"]).value
            if not mrd:
                continue
            counts["Total Patients"] += 1
            status = self.ws.cell(row=r, column=COL_INDEX["Status"]).value
            if status in _ALL_STATUSES:
                counts[status] += 1
        return counts

    def all_rows(self) -> list:
        """Every patient row in the workbook as a dict keyed by column name."""
        rows = []
        for r in range(2, self.ws.max_row + 1):
            if not self.ws.cell(row=r, column=COL_INDEX["MRD Number"]).value:
                continue
            rows.append({name: self.ws.cell(row=r, column=i).value for name, i in COL_INDEX.items()})
        return rows

    def summary(self) -> dict:
        """Whole-workbook overview: status counts plus money totals.
        Billed / difference totals only cover patients that have a bill."""
        rows = self.all_rows()
        counts = self.overall_counts()

        def _num(v):
            return v if isinstance(v, (int, float)) else 0

        billed = [r for r in rows if r["Billed Amount"] is not None]
        return {
            "counts": counts,
            "rows": rows,
            "total_estimate": sum(_num(r["Estimate Amount"]) for r in rows),
            "total_billed": sum(_num(r["Billed Amount"]) for r in billed),
            "billed_estimate": sum(_num(r["Estimate Amount"]) for r in billed),
            "net_difference": sum(_num(r["Difference"]) for r in billed),
        }
